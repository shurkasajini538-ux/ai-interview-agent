# -*- coding: utf-8 -*-
"""生成项目计划书 Word 版"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(11)

t = doc.add_heading("AI面试陪练系统（基于 openJiuwen）项目计划书", level=1)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER

sections = [
    ("一、项目背景与痛点", [
        "大学生求职、保研准备面试时，普遍存在“有项目但不会讲”的问题：项目经历缺少量化结果，技术细节解释不清，遇到追问就卡壳。",
        "通用大模型只能给通用建议，无法读取个人简历做针对性追问、打分和改写，练完还是不会答。",
    ]),
    ("二、解决方案", [
        "做一个 AI 面试官：上传简历后自动解析教育背景、项目经历、技能点，按固定流程陪练：简历解析 → 问题生成（≥5题） → 学生回答 → 追问 → 打分 → STAR 优化回答。",
        "基于 openJiuwen 框架实现：Knowledge 模块存简历和项目材料；Agent Workflow 编排六步流程；Memory 记录“缺少量化结果、技术细节不清”等薄弱点，越练越准。",
    ]),
    ("三、目标群体", [
        "1. 准备求职、保研的本科生，尤其是有项目但缺表达指导的中腰部高校学生；",
        "2. 高校就业辅导老师，可作为课堂辅助工具批量辅导。",
    ]),
    ("四、项目远景", [
        "让每个学生都拥有 24 小时在线的个人面试教练，把散落的面经与教师经验沉淀为高校可复用的面试知识库，缩小信息差。",
    ]),
    ("五、竞争分析", [
        "通用大模型（ChatGPT、DeepSeek）：只能通用问答，无个人简历追问、无打分、无记忆；",
        "面经平台（牛客、实习僧）：只有题库，无个性化陪练；",
        "线下付费 1 对 1：贵且不可规模化。",
        "本项目差异化：个人简历 RAG + 可控 Workflow + 长期记忆，形成练测改闭环。",
    ]),
    ("六、功能与答题指标对应", [
        "简历片段输入 → config/resume_sample.md；生成≥5题、≥1轮追问、1次评分、优化回答 → python demo.py 一键演示；",
        "Agent 工作流图 → docs/workflow.mmd；普通大模型与 openJiuwen Agent 区别 → README 对比表；演示页 → streamlit run app.py。",
    ]),
    ("七、进度计划（4周）", [
        "第1周：跑通 Demo（已完成，代码已开源）；",
        "第2周：接入 Studio 可视化流程与知识库；",
        "第3周：完成网页演示版，录制 2 分钟演示视频；",
        "第4周：完成答辩 PPT，在 GitCode 提交报名。",
    ]),
    ("八、团队分工", [
        "后端1人：Agent 流程与评分逻辑；前端1人：演示页与视频录制；材料1人：PPT 与计划书。单人可全包，多人更快。",
    ]),
    ("九、预期成果", [
        "可运行 Demo 一套、GitCode/GitHub 开源仓库、2 分钟演示视频、答辩 PPT 与本计划书。",
    ]),
]

for title, paras in sections:
    doc.add_heading(title, level=2)
    for p in paras:
        doc.add_paragraph(p)

doc.add_paragraph("")
p = doc.add_paragraph("仓库地址：")
p.add_run("https://gitcode.com/easyyyyyyyy/ai-interview-agent").font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

out = r"C:\Users\47148\Desktop\ai-interview-agent\docs\项目计划书.docx"
doc.save(out)
print("saved:", out)
